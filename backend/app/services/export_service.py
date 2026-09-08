"""Supplier-ready purchase order exports (Excel / PDF)."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.schemas.inventory import ExportRequest

HEADERS = ["Product ID", "Store", "Family", "Order Qty", "Unit Cost", "Line Total"]


def _rows(payload: ExportRequest) -> list[list[str | float]]:
    rows: list[list[str | float]] = []
    for line in payload.lines:
        unit_cost = line.unit_cost if line.unit_cost is not None else 0.0
        rows.append(
            [
                line.product_id,
                line.store_nbr,
                line.family,
                round(line.order_qty, 2),
                round(unit_cost, 2),
                round(line.order_qty * unit_cost, 2),
            ]
        )
    return rows


def build_excel(payload: ExportRequest) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Purchase Order"

    sheet.append([f"ForecastIQ Purchase Order - {payload.supplier_name}"])
    sheet.append([f"Issued: {date.today().isoformat()}", f"Mode: {payload.confidence_mode.value}"])
    sheet.append([])
    sheet.append(HEADERS)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0D9488")
    for cell in sheet[4]:
        cell.font = header_font
        cell.fill = header_fill

    rows = _rows(payload)
    for row in rows:
        sheet.append(row)
    sheet.append([])
    sheet.append(["", "", "TOTAL", sum(row[3] for row in rows), "", sum(row[5] for row in rows)])

    for column, width in zip("ABCDEF", (16, 8, 28, 12, 12, 14)):
        sheet.column_dimensions[column].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_pdf(payload: ExportRequest) -> bytes:
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, title="ForecastIQ Purchase Order")
    styles = getSampleStyleSheet()
    rows = _rows(payload)

    elements = [
        Paragraph("ForecastIQ Purchase Order", styles["Title"]),
        Paragraph(f"Supplier: {payload.supplier_name}", styles["Normal"]),
        Paragraph(f"Issued: {date.today().isoformat()}", styles["Normal"]),
        Paragraph(f"Confidence mode: {payload.confidence_mode.value}", styles["Normal"]),
        Spacer(1, 16),
    ]

    table_data = [HEADERS, *[[str(value) for value in row] for row in rows]]
    table_data.append(
        ["", "", "TOTAL", str(round(sum(row[3] for row in rows), 2)), "", str(round(sum(row[5] for row in rows), 2))]
    )
    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D9488")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#1B3A5C")),
                ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
            ]
        )
    )
    elements.append(table)
    document.build(elements)
    return buffer.getvalue()
